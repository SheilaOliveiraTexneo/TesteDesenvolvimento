from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import login_required, current_user
from app import db
from app.models import Solicitacao, Partida, Produto, Solicitante, AuditoriaLog
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

solicitacoes_bp = Blueprint('solicitacoes', __name__, url_prefix='/solicitacoes')

@solicitacoes_bp.route('/')
@login_required
def listar():
    """Lista todas as solicitações"""
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '')
    status = request.args.get('status', '')
    prioridade = request.args.get('prioridade', '')
    sort_by = request.args.get('sort', 'id_solicitacao')
    order = request.args.get('order', 'desc')
    
    query = Solicitacao.query
    
    if search:
        query = query.join(Partida).join(Produto).join(Solicitante).filter(
            db.or_(
                Solicitacao.id_solicitacao.ilike(f'%{search}%'),
                Produto.codigo.ilike(f'%{search}%'),
                Solicitante.codigo.ilike(f'%{search}%')
            )
        )
    
    if status:
        query = query.filter_by(status=status)
    
    if prioridade:
        query = query.filter_by(prioridade=int(prioridade))
    
    # Ordenamento
    order_column = getattr(Solicitacao, sort_by, Solicitacao.id_solicitacao)
    query = query.order_by(order_column.asc() if order == 'asc' else order_column.desc())
    
    paginated = query.paginate(page=page, per_page=10)
    
    return render_template(
        'solicitacoes/listar.html',
        solicitacoes=paginated.items,
        paginated=paginated,
        search=search,
        status=status,
        prioridade=prioridade,
        sort_by=sort_by,
        order=order
    )

@solicitacoes_bp.route('/nova', methods=['GET', 'POST'])
@login_required
def nova():
    """Criar nova solicitação"""
    if not current_user.tem_permissao('criar'):
        flash('Você não tem permissão para criar solicitações', 'danger')
        return redirect(url_for('solicitacoes.listar'))
    
    if request.method == 'POST':
        id_partida = request.form.get('id_partida').strip()
        codigo_solicitante = request.form.get('codigo_solicitante').strip()
        prioridade = int(request.form.get('prioridade', 3))
        lote = request.form.get('lote', '').strip()
        numero_rocas = request.form.get('numero_rocas')
        observacoes = request.form.get('observacoes', '').strip()
        
        # Buscar partida
        partida = Partida.query.filter_by(id_partida=id_partida).first()
        if not partida:
            flash(f'Partida {id_partida} não encontrada', 'danger')
            return redirect(url_for('solicitacoes.nova'))
        
        # Buscar solicitante
        solicitante = Solicitante.query.filter_by(codigo=codigo_solicitante).first()
        if not solicitante:
            flash(f'Solicitante {codigo_solicitante} não encontrado', 'danger')
            return redirect(url_for('solicitacoes.nova'))
        
        # Gerar ID da solicitação
        id_solicitacao = Solicitacao.gerar_id_solicitacao()
        
        solicitacao = Solicitacao(
            id_solicitacao=id_solicitacao,
            partida_id=partida.id,
            produto_id=partida.produto_id,
            solicitante_id=solicitante.id,
            prioridade=prioridade,
            lote=lote,
            numero_rocas=int(numero_rocas) if numero_rocas else None,
            status='solicitado',
            observacoes=observacoes
        )
        
        db.session.add(solicitacao)
        db.session.flush()
        
        # Registrar auditoria
        AuditoriaLog.registrar(
            'solicitacoes',
            'INSERT',
            solicitacao.id,
            dados_depois=str(solicitacao.to_dict()),
            descricao=f'Nova solicitação criada: {id_solicitacao}'
        )
        
        db.session.commit()
        flash(f'Solicitação {id_solicitacao} criada com sucesso', 'success')
        return redirect(url_for('solicitacoes.listar'))
    
    return render_template('solicitacoes/form.html', solicitacao=None)

@solicitacoes_bp.route('/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def editar(id):
    """Editar solicitação"""
    if not current_user.tem_permissao('editar'):
        flash('Você não tem permissão para editar solicitações', 'danger')
        return redirect(url_for('solicitacoes.listar'))
    
    solicitacao = Solicitacao.query.get_or_404(id)
    
    if request.method == 'POST':
        prioridade = int(request.form.get('prioridade', 3))
        status = request.form.get('status', 'solicitado')
        lote = request.form.get('lote', '').strip()
        numero_rocas = request.form.get('numero_rocas')
        observacoes = request.form.get('observacoes', '').strip()
        
        dados_antes = str(solicitacao.to_dict())
        
        solicitacao.prioridade = prioridade
        solicitacao.status = status
        solicitacao.lote = lote
        solicitacao.numero_rocas = int(numero_rocas) if numero_rocas else None
        solicitacao.observacoes = observacoes
        
        db.session.commit()
        
        # Registrar auditoria
        AuditoriaLog.registrar(
            'solicitacoes',
            'UPDATE',
            solicitacao.id,
            dados_antes=dados_antes,
            dados_depois=str(solicitacao.to_dict()),
            descricao=f'Solicitação {solicitacao.id_solicitacao} atualizada'
        )
        
        flash('Solicitação atualizada com sucesso', 'success')
        return redirect(url_for('solicitacoes.listar'))
    
    return render_template('solicitacoes/form.html', solicitacao=solicitacao)

@solicitacoes_bp.route('/<int:id>/deletar', methods=['POST'])
@login_required
def deletar(id):
    """Deletar solicitação"""
    if not current_user.tem_permissao('deletar_solicitacao'):
        flash('Você não tem permissão para deletar solicitações', 'danger')
        return redirect(url_for('solicitacoes.listar'))
    
    solicitacao = Solicitacao.query.get_or_404(id)
    
    if solicitacao.status == 'finalizado':
        flash('Não é possível deletar uma solicitação finalizada', 'danger')
        return redirect(url_for('solicitacoes.listar'))
    
    dados_antes = str(solicitacao.to_dict())
    id_solicitacao = solicitacao.id_solicitacao
    
    db.session.delete(solicitacao)
    db.session.commit()
    
    # Registrar auditoria
    AuditoriaLog.registrar(
        'solicitacoes',
        'DELETE',
        id,
        dados_antes=dados_antes,
        descricao=f'Solicitação deletada: {id_solicitacao}'
    )
    
    flash('Solicitação deletada com sucesso', 'success')
    return redirect(url_for('solicitacoes.listar'))

@solicitacoes_bp.route('/exportar/excel')
@login_required
def exportar_excel():
    """Exportar solicitações para Excel"""
    solicitacoes = Solicitacao.query.all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = 'Solicitações'
    
    # Títulos
    headers = [
        'ID Solicitação',
        'Partida',
        'Produto',
        'Solicitante',
        'Prioridade',
        'Lote',
        'Nº Rocas',
        'Data',
        'Status',
        'Observações'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    
    # Dados
    for row, sol in enumerate(solicitacoes, 2):
        ws.cell(row=row, column=1, value=sol.id_solicitacao)
        ws.cell(row=row, column=2, value=sol.partida.id_partida)
        ws.cell(row=row, column=3, value=f"{sol.produto.codigo} - {sol.produto.descricao}")
        ws.cell(row=row, column=4, value=sol.solicitante.nome)
        ws.cell(row=row, column=5, value=sol.prioridade)
        ws.cell(row=row, column=6, value=sol.lote or '')
        ws.cell(row=row, column=7, value=sol.numero_rocas or '')
        ws.cell(row=row, column=8, value=sol.data_solicitacao.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row, column=9, value=sol.status)
        ws.cell(row=row, column=10, value=sol.observacoes or '')
    
    # Ajustar largura das colunas
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 20
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 25
    
    # Salvar em memória
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'solicitacoes_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    )
